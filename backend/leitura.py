"""Leitura dos arquivos de extrato: OFX, CSV e Excel (.xlsx).

Cada leitor devolve uma lista de lançamentos no formato:
    {"data": "2026-08-01", "descricao": "PIX ENVIADO ...", "valor": -120.5}
Valores negativos são saídas; positivos, entradas.
"""
import csv
import io
import re
from datetime import date, datetime

from categorizacao import normalizar


class ErroLeitura(Exception):
    """Erro com mensagem pronta para mostrar ao cliente."""


# ---------- valores e datas ----------

_NUMERO_INICIAL = re.compile(r"^\d*\.?\d*")
_MILHAR = re.compile(r"^\d{1,3}(\.\d{3})+$")


def ler_valor(bruto):
    """Converte '1.234,56', '-R$ 10,00', '(50,00)', '120,00 D' etc. em número.
    Devolve None se não for um valor."""
    if bruto is None:
        return None
    if isinstance(bruto, bool):
        return None
    if isinstance(bruto, (int, float)):
        return float(bruto)
    s = str(bruto).strip()
    if not s:
        return None
    negativo = bool(re.match(r"^\(.*\)$", s)) or "-" in s or bool(re.search(r"\sD$", s, re.I))
    s = re.sub(r"[^\d.,]", "", s)
    if not s:
        return None
    virgula, ponto = s.rfind(","), s.rfind(".")
    if virgula > -1 and ponto > -1:
        s = s.replace(".", "").replace(",", ".", 1) if virgula > ponto else s.replace(",", "")
    elif virgula > -1:
        s = s.replace(".", "").replace(",", ".", 1)
    elif _MILHAR.match(s):
        s = s.replace(".", "")
    m = _NUMERO_INICIAL.match(s).group(0)
    if m in ("", "."):
        return None
    v = float(m)
    return -abs(v) if negativo else v


def ler_data(bruto):
    """Aceita 31/08/2026, 31-08-26, 2026-08-31, 20260831 ou uma data do Excel."""
    if isinstance(bruto, datetime):
        return bruto.date()
    if isinstance(bruto, date):
        return bruto
    s = str(bruto or "").strip()
    m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", s)
    if m:
        ano = int(m.group(3))
        if ano < 100:
            ano += 2000
        return _montar_data(ano, int(m.group(2)), int(m.group(1)))
    m = re.match(r"^(\d{4})-?(\d{2})-?(\d{2})", s)
    if m:
        return _montar_data(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _montar_data(ano, mes, dia):
    try:
        return date(ano, mes, dia)
    except ValueError:
        return None


# ---------- tabelas (CSV e Excel) ----------

def _celula(linha, i):
    if i < 0 or i >= len(linha):
        return None
    return linha[i]


def ler_tabela(linhas):
    """Lê uma tabela (lista de linhas, cada linha uma lista de células)
    descobrindo sozinho as colunas de data, descrição e valor."""
    linhas = [l for l in linhas if any(str(c if c is not None else "").strip() for c in l)]
    if not linhas:
        raise ErroLeitura("O arquivo está vazio.")

    # procura a linha de cabeçalho (a que tem uma coluna "data")
    cab_idx, cab = -1, None
    for i, linha in enumerate(linhas[:25]):
        celulas = [normalizar(c if c is not None else "") for c in linha]
        if any(" data " in c or " date " in c for c in celulas):
            cab_idx, cab = i, celulas
            break

    def coluna(chaves, ignorar=-1):
        if not cab:
            return -1
        for chave in chaves:
            for k, titulo in enumerate(cab):
                if k != ignorar and chave in titulo:
                    return k
        return -1

    c_data = coluna([" data ", " date "])
    c_desc = coluna(["descri", "historico", "lancamento", "memo", "detalhe", "estabelecimento", "titulo"], c_data)
    c_valor = coluna([" valor", "montante", "quantia", "amount"], c_data)
    c_cred = coluna(["credito", "entrada"], c_data)
    c_deb = coluna(["debito", "saida"], c_data)
    c_tipo = coluna([" tipo ", " d c ", "natureza"], c_data)
    inicio = cab_idx + 1
    if not cab:
        c_data, inicio = 0, 0

    saida = []
    for linha in linhas[inicio:]:
        dt = ler_data(_celula(linha, c_data))
        if not dt:
            continue
        desc = _celula(linha, c_desc) if c_desc > -1 else None
        desc = str(desc).strip() if desc not in (None, "") else ""
        if not desc:
            candidatas = [
                str(x).strip() for ix, x in enumerate(linha)
                if ix != c_data and x not in (None, "") and ler_valor(x) is None
            ]
            candidatas.sort(key=len, reverse=True)
            desc = candidatas[0] if candidatas else "Sem descrição"
        if re.search("saldo", desc, re.I):
            continue

        v = ler_valor(_celula(linha, c_valor)) if c_valor > -1 else None
        if v is None and (c_cred > -1 or c_deb > -1):
            cr = ler_valor(_celula(linha, c_cred)) if c_cred > -1 else None
            db = ler_valor(_celula(linha, c_deb)) if c_deb > -1 else None
            if db is not None and db != 0:
                v = -abs(db)
            elif cr is not None:
                v = abs(cr)
        if v is None and not cab:
            for k in range(len(linha) - 1, 0, -1):
                t = ler_valor(linha[k])
                if t is not None:
                    v = t
                    break
        if v is None or v == 0:
            continue
        tipo = _celula(linha, c_tipo) if c_tipo > -1 else None
        if v > 0 and tipo is not None and re.match(r"^\s*(d|deb)", str(tipo), re.I):
            v = -v
        saida.append({"data": dt.isoformat(), "descricao": desc, "valor": round(v, 2)})

    if not saida:
        raise ErroLeitura("Não encontrei lançamentos neste arquivo. Confira se ele tem colunas de data, descrição e valor.")
    return saida


def ler_csv(texto):
    linhas_texto = [l for l in texto.splitlines() if l.strip()]
    if not linhas_texto:
        raise ErroLeitura("O arquivo está vazio.")
    amostra = "\n".join(linhas_texto[:20])
    separador = sorted([";", "\t", ","], key=lambda d: -amostra.count(d))[0]
    leitor = csv.reader(io.StringIO("\n".join(linhas_texto)), delimiter=separador, quotechar='"')
    return ler_tabela([[c.strip() for c in linha] for linha in leitor])


def ler_ofx(texto):
    blocos = re.split(r"<STMTTRN>", texto, flags=re.I)[1:]

    def tag(bloco, nome):
        m = re.search("<" + nome + r">([^<\r\n]*)", bloco, re.I)
        return m.group(1).strip() if m else ""

    saida = []
    for b in blocos:
        dt = ler_data(tag(b, "DTPOSTED")[:8])
        bruto = tag(b, "TRNAMT").replace(",", ".", 1)
        try:
            v = float(re.match(r"^[-+]?\d*\.?\d*", bruto).group(0))
        except ValueError:
            v = None
        desc = tag(b, "MEMO") or tag(b, "NAME") or "Sem descrição"
        if dt and v:
            saida.append({"data": dt.isoformat(), "descricao": desc, "valor": round(v, 2)})
    if not saida:
        raise ErroLeitura("Não encontrei lançamentos neste arquivo OFX.")
    return saida


def ler_xlsx(conteudo):
    try:
        from openpyxl import load_workbook
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:
        raise ErroLeitura("Não consegui abrir a planilha Excel. Tente salvar de novo como .xlsx ou CSV.")
    todas, primeiro_cab = [], None
    for aba in livro.worksheets:
        linhas = [
            ["" if c is None else c for c in linha]
            for linha in aba.iter_rows(values_only=True)
            if any(c not in (None, "") for c in linha)
        ]
        if not linhas:
            continue
        if primeiro_cab is None:
            primeiro_cab = linhas[0]
            todas.extend(linhas)
        else:
            todas.extend(linhas[1:] if linhas[0] == primeiro_cab else linhas)
    if not todas:
        raise ErroLeitura("A planilha Excel está vazia.")
    return ler_tabela(todas)


def decodificar(conteudo):
    try:
        return conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        return conteudo.decode("latin-1")


def ler_arquivo(nome, conteudo):
    """Escolhe o leitor pelo tipo do arquivo e devolve os lançamentos
    em ordem de data."""
    nome_min = (nome or "").lower()
    if nome_min.endswith(".pdf"):
        raise ErroLeitura("PDF ainda não é aceito. Exporte em OFX ou CSV.")
    if nome_min.endswith((".xlsx", ".xlsm")):
        lista = ler_xlsx(conteudo)
    elif nome_min.endswith((".xls", ".ods")):
        raise ErroLeitura("Este formato de planilha não é aceito. Salve como .xlsx ou CSV.")
    else:
        texto = decodificar(conteudo)
        lista = ler_ofx(texto) if re.search(r"<OFX>|<STMTTRN>", texto, re.I) else ler_csv(texto)
    lista.sort(key=lambda l: l["data"])
    return lista
